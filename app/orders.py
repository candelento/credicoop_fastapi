import base64
import hashlib
import io
import json
import re
import unicodedata
from dataclasses import dataclass
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path
from urllib.parse import quote
from uuid import NAMESPACE_URL, uuid5

import httpx
from cryptography.hazmat.primitives import hashes, serialization
from cryptography.hazmat.primitives.asymmetric import padding
from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException
from pypdf import PdfReader
from pypdf.errors import PdfReadError

from .config import Settings
from .journal import Journal


class DriveOrderError(Exception):
    pass


class SupplierMasterError(Exception):
    pass


class OrderParseError(Exception):
    pass


@dataclass(frozen=True)
class Supplier:
    name: str
    document: str
    cbu: str


AMOUNT_PATTERN = re.compile(r'(?<![\d.,+\-])(?:\d{1,3}(?:[. ]\d{3})+|\d+)(?:[,.]\d{2})(?![\d.,])')
METHOD_PATTERN = re.compile(r'(?<![A-Z0-9])(TRB|CHE)(?![A-Z0-9])', re.IGNORECASE)
DATE_PATTERN = re.compile(r'(?<!\d)(\d{2}/\d{2}/(?:\d{2}|\d{4}))(?!\d)')


def normalize_text(value: str) -> str:
    decomposed = unicodedata.normalize('NFKD', value)
    plain = ''.join(
        character for character in decomposed
        if not unicodedata.combining(character) and (character.isalnum() or character.isspace())
    )
    return ' '.join(plain.upper().split())


def _digits(value: object) -> str:
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return ''.join(character for character in str(value or '') if character.isdigit())


def load_supplier_master(path: Path) -> dict[str, list[Supplier]]:
    if not path.is_file():
        raise SupplierMasterError(f'No se encuentra el maestro de beneficiarios configurado: {path.name}.')
    workbook = None
    try:
        workbook = load_workbook(path, read_only=True, data_only=True)
        sheet = workbook['Beneficiarios Transferencias'] if 'Beneficiarios Transferencias' in workbook.sheetnames else workbook.active
        rows = sheet.iter_rows(values_only=True)
        headers = None
        for row_number, row in enumerate(rows):
            candidate = {normalize_text(str(value or '')).replace(' ', ''): index
                         for index, value in enumerate(row) if value is not None}
            name_column = next((candidate[key] for key in ('NOMBRE', 'PROVEEDOR', 'NOMBREBENEFICIARIO')
                                if key in candidate), None)
            document_column = next((candidate[key] for key in ('CUITCUILCDI', 'CUIT', 'CUIL', 'DOCUMENTO')
                                    if key in candidate), None)
            cbu_column = candidate.get('CBU')
            if name_column is not None and document_column is not None and cbu_column is not None:
                headers = (name_column, document_column, cbu_column)
                break
            if row_number >= 9:
                break
        if headers is None:
            raise SupplierMasterError('El maestro no contiene columnas reconocibles de nombre, CUIT/CUIL y CBU.')

        name_index, document_index, cbu_index = headers
        suppliers: dict[str, list[Supplier]] = {}
        for row in rows:
            values = list(row)
            if max(headers) >= len(values) or not values[name_index]:
                continue
            name = ' '.join(str(values[name_index]).split())
            key = normalize_text(name)
            if len(key) < 3:
                continue
            supplier = Supplier(name=name, document=_digits(values[document_index]), cbu=_digits(values[cbu_index]))
            if supplier not in suppliers.setdefault(key, []):
                suppliers[key].append(supplier)
        if not suppliers:
            raise SupplierMasterError('El maestro de beneficiarios no contiene proveedores.')
        return suppliers
    except (InvalidFileException, OSError, ValueError, KeyError) as exc:
        raise SupplierMasterError('No se pudo leer el maestro de beneficiarios.') from exc
    finally:
        if workbook is not None:
            workbook.close()


def merge_supplier_sources(*sources: dict[str, list[Supplier]]) -> dict[str, list[Supplier]]:
    merged: dict[str, list[Supplier]] = {}
    for source in sources:
        for key, items in source.items():
            bucket = merged.setdefault(key, [])
            for supplier in items:
                if supplier not in bucket:
                    bucket.append(supplier)
    return merged


def load_homologation_suppliers(path: Path) -> dict[str, list[Supplier]]:
    if not path.is_file():
        return {}
    try:
        payload = json.loads(path.read_text(encoding='utf-8'))
        beneficiaries = payload.get('beneficiarios', []) if isinstance(payload, dict) else []
    except (OSError, ValueError, TypeError) as exc:
        raise SupplierMasterError('No se pudo leer la lista de beneficiarios de homologación.') from exc

    suppliers: dict[str, list[Supplier]] = {}
    for beneficiary in beneficiaries:
        if not isinstance(beneficiary, dict):
            continue
        if beneficiary.get('uso') != 'emision' or beneficiary.get('tipo') != 'transferencia':
            continue
        name = ' '.join(str(beneficiary.get('nombre') or '').split())
        key = normalize_text(name)
        if len(key) < 3:
            continue
        supplier = Supplier(
            name=name,
            document=_digits(beneficiary.get('documento')),
            cbu=_digits(beneficiary.get('cbuCvu')),
        )
        if supplier not in suppliers.setdefault(key, []):
            suppliers[key].append(supplier)
    return suppliers


def _amount_value(token: str) -> str:
    compact = token.replace(' ', '')
    if ',' in compact:
        compact = compact.replace('.', '').replace(',', '.')
    elif compact.count('.') > 1:
        raise OrderParseError('El formato del importe no es claro.')
    try:
        amount_decimal = Decimal(compact)
        if not amount_decimal.is_finite() or amount_decimal <= 0 or amount_decimal != amount_decimal.quantize(Decimal('.01')):
            raise OrderParseError('El importe no es un monto decimal positivo válido.')
        amount = format(amount_decimal, '.2f')
    except (InvalidOperation, ValueError):
        raise OrderParseError('No se pudo interpretar el importe de la orden.') from None
    return amount


def _date_value(token: str) -> date:
    for format_string in ('%d/%m/%Y', '%d/%m/%y'):
        try:
            return datetime.strptime(token, format_string).date()
        except ValueError:
            continue
    raise OrderParseError('La fecha indicada en la línea CHE no es válida.')


def order_origin(file_id: str) -> str:
    return str(uuid5(NAMESPACE_URL, f'credicoop-op-drive:{file_id}'))


def parse_order(filename: str, file_id: str, content: bytes,
                suppliers: dict[str, list[Supplier]], operational_today: date) -> dict:
    if not filename.lower().startswith('op-') or not filename.lower().endswith('.pdf'):
        raise OrderParseError('El archivo no tiene el patrón op-*.pdf.')
    if not content.startswith(b'%PDF'):
        raise OrderParseError('El archivo descargado no es un PDF válido.')
    try:
        reader = PdfReader(io.BytesIO(content))
        if reader.is_encrypted:
            raise OrderParseError('El PDF está cifrado y requiere revisión manual.')
        lines = '\n'.join(page.extract_text() or '' for page in reader.pages).splitlines()
    except (PdfReadError, OSError, ValueError) as exc:
        raise OrderParseError('No se pudo extraer texto del PDF; requiere revisión manual.') from exc

    method_lines = [(line, METHOD_PATTERN.search(line)) for line in lines]
    method_lines = [(line, match) for line, match in method_lines if match]
    order = {
        'source_id': file_id,
        'source_name': filename,
        'id_origen': order_origin(file_id),
        'sha256': hashlib.sha256(content).hexdigest(),
        'status': 'NEEDS_REVIEW',
        'reason': None,
        'method': None,
        'supplier': None,
        'amount': None,
        'cheque_date': None,
        'cheque_type': None,
    }
    if len(method_lines) != 1:
        order['reason'] = 'No se encontró una única línea con TRB o CHE.'
        return order

    line, match = method_lines[0]
    method = match.group(1).upper()
    order['method'] = method
    line_after_method = line[match.end():]
    amounts = AMOUNT_PATTERN.findall(line_after_method)
    if len(amounts) != 1:
        order['reason'] = 'La línea de TRB/CHE no contiene un único importe reconocible.'
        return order
    try:
        order['amount'] = _amount_value(amounts[0])
    except OrderParseError as exc:
        order['reason'] = str(exc)
        return order
    if Decimal(order['amount']) >= Decimal('10000000000000'):
        order['reason'] = 'El importe supera el máximo admitido para una instrucción bancaria.'
        return order

    matched: dict[str, list[Supplier]] = {}
    for line_text in lines:
        normalized_line = f' {normalize_text(line_text)} '
        for supplier_name, entries in suppliers.items():
            if len(supplier_name) >= 6 and f' {supplier_name} ' in normalized_line:
                matched[supplier_name] = entries
    unique_suppliers = {entry for entries in matched.values() for entry in entries}
    if len(unique_suppliers) != 1:
        order['reason'] = ('No hay una coincidencia única del proveedor en el maestro.'
                           if not unique_suppliers else 'El PDF coincide con más de un beneficiario del maestro.')
        return order
    supplier = next(iter(unique_suppliers))
    if len(supplier.document) != 11 or len(supplier.cbu) != 22:
        order['supplier'] = {'name': supplier.name, 'document': supplier.document, 'cbu': supplier.cbu}
        order['reason'] = 'El CUIT/CUIL o CBU del proveedor no tiene el formato requerido.'
        return order
    order['supplier'] = {'name': supplier.name, 'document': supplier.document, 'cbu': supplier.cbu}

    if method == 'CHE':
        dates = DATE_PATTERN.findall(line)
        if len(dates) != 1:
            order['reason'] = 'La línea CHE no contiene una única fecha de pago reconocible.'
            return order
        try:
            cheque_date = _date_value(dates[0])
        except OrderParseError as exc:
            order['reason'] = str(exc)
            return order
        if cheque_date < operational_today:
            order['reason'] = 'La fecha del eCheq es anterior a la fecha operativa; requiere revisión.'
            return order
        order['cheque_date'] = cheque_date.isoformat()
        order['cheque_type'] = 'ECHC' if cheque_date == operational_today else 'ECHD'

    order['status'] = 'READY'
    return order


def build_order_instruction(order: dict, debit_cbu: str, operational_today: date) -> dict:
    supplier = order['supplier']
    safe_name = normalize_text(supplier['name'])
    if not safe_name or len(safe_name) > 100:
        raise OrderParseError('El nombre del proveedor no es válido para la instrucción bancaria.')
    if order['method'] == 'TRB':
        return {
            'idOrigen': order['id_origen'],
            'cbuCuentaDebito': debit_cbu,
            'beneficiarios': [{
                'cbuCvu': supplier['cbu'],
                'monto': order['amount'],
                'moneda': 'ARS',
                'concepto': 'VAR',
                'cui': supplier['document'],
                'nombre': safe_name,
                'observaciones': 'ORDEN DE PAGO',
            }],
        }
    cheque_date = date.fromisoformat(order['cheque_date'])
    cheque_type = 'ECHC' if cheque_date == operational_today else 'ECHD'
    if cheque_type != order['cheque_type']:
        raise OrderParseError('La fecha del eCheq difiere de la fecha operativa; sincronizá de nuevo para revisar.')
    return {
        'idOrigen': order['id_origen'],
        'cbuCuentaDebito': debit_cbu,
        'echeqs': [{
            'monto': order['amount'],
            'fechaPago': order['cheque_date'],
            'motivoPago': 'ORDEN DE PAGO',
            'caracter': '1',
            'beneficiarioNombre': safe_name,
            'beneficiarioDocumentoTipo': 'CUIT',
            'beneficiarioDocumento': supplier['document'],
            'concepto': 'VAR',
            'tipoCheque': cheque_type,
        }],
    }


class DriveOrderService:
    def __init__(self, settings: Settings, project_root: Path):
        self.settings = settings
        self.project_root = project_root
        self.credentials_path = self._resolve_path(settings.drive_credentials_path)
        self.master_path = self._resolve_path(settings.supplier_master_path)

    def _resolve_path(self, path: Path) -> Path:
        return path if path.is_absolute() else self.project_root / path

    @staticmethod
    def _encode(value: bytes) -> str:
        return base64.urlsafe_b64encode(value).rstrip(b'=').decode('ascii')

    def _signed_assertion(self) -> tuple[str, str]:
        try:
            credentials = json.loads(self.credentials_path.read_text(encoding='utf-8'))
            if not isinstance(credentials, dict):
                raise DriveOrderError('El archivo de credenciales de Drive no tiene un objeto JSON válido.')
            if credentials.get('type') != 'service_account':
                raise DriveOrderError('La credencial de Drive configurada no es una cuenta de servicio.')
            if credentials.get('token_uri') != 'https://oauth2.googleapis.com/token':
                raise DriveOrderError('La credencial de Drive tiene un endpoint de autenticación no permitido.')
            if not isinstance(credentials.get('client_email'), str) or not isinstance(credentials.get('private_key'), str):
                raise DriveOrderError('La credencial de Drive no contiene los campos requeridos de cuenta de servicio.')
            now = int(datetime.now().timestamp())
            header = self._encode(b'{"alg":"RS256","typ":"JWT"}')
            claims = self._encode(json.dumps({
                'iss': credentials['client_email'],
                'scope': 'https://www.googleapis.com/auth/drive.readonly',
                'aud': credentials['token_uri'],
                'iat': now,
                'exp': now + 300,
            }, separators=(',', ':')).encode('utf-8'))
            unsigned = f'{header}.{claims}'.encode('ascii')
            private_key = serialization.load_pem_private_key(credentials['private_key'].encode(), password=None)
            assertion = f'{unsigned.decode("ascii")}.{self._encode(private_key.sign(unsigned, padding.PKCS1v15(), hashes.SHA256()))}'
            return credentials['token_uri'], assertion
        except DriveOrderError:
            raise
        except (OSError, ValueError, KeyError, TypeError) as exc:
            raise DriveOrderError('No se pudo cargar la credencial RSA de solo lectura de Drive.') from exc

    async def _access_token(self, client: httpx.AsyncClient) -> str:
        token_url, assertion = self._signed_assertion()
        try:
            response = await client.post(token_url, data={
                'grant_type': 'urn:ietf:params:oauth:grant-type:jwt-bearer',
                'assertion': assertion,
            })
            response.raise_for_status()
            payload = response.json()
            token = payload.get('access_token') if isinstance(payload, dict) else None
            if not isinstance(token, str) or not token:
                raise DriveOrderError('Google Drive no devolvió un token de acceso.')
            return token
        except DriveOrderError:
            raise
        except (httpx.HTTPError, ValueError) as exc:
            raise DriveOrderError('No se pudo autenticar con Google Drive en modo de solo lectura.') from exc

    async def _list_files(self, client: httpx.AsyncClient, token: str) -> list[dict]:
        folder_id = self.settings.drive_folder_id
        if not re.fullmatch(r'[A-Za-z0-9_-]{10,200}', folder_id):
            raise DriveOrderError('El ID de carpeta de Google Drive no es válido.')
        headers = {'Authorization': f'Bearer {token}'}
        files: list[dict] = []
        page_token = None
        try:
            while True:
                params = {
                    'q': f"'{folder_id}' in parents and trashed = false and mimeType = 'application/pdf'",
                    'pageSize': 1000,
                    'orderBy': 'name',
                    'fields': 'nextPageToken,files(id,name,mimeType)',
                    'supportsAllDrives': 'true',
                    'includeItemsFromAllDrives': 'true',
                }
                if page_token:
                    params['pageToken'] = page_token
                response = await client.get('https://www.googleapis.com/drive/v3/files', params=params, headers=headers)
                response.raise_for_status()
                payload = response.json()
                if not isinstance(payload, dict) or not isinstance(payload.get('files', []), list):
                    raise DriveOrderError('Google Drive devolvió un listado de archivos con formato inesperado.')
                files.extend(
                    item for item in payload.get('files', [])
                    if item.get('mimeType') == 'application/pdf'
                    and item.get('name', '').lower().startswith('op-')
                    and item.get('name', '').lower().endswith('.pdf')
                    and item.get('id')
                )
                page_token = payload.get('nextPageToken')
                if not page_token:
                    return files
        except (httpx.HTTPError, ValueError) as exc:
            raise DriveOrderError('No se pudo listar la carpeta configurada de Google Drive.') from exc

    async def sync(self, journal: Journal) -> dict:
        suppliers = load_supplier_master(self.master_path)
        if self.settings.is_homologation:
            suppliers = merge_supplier_sources(
                suppliers,
                load_homologation_suppliers(self.project_root / 'ejemplos' / 'beneficiarios_homologacion.json'),
            )
        async with httpx.AsyncClient(timeout=30) as client:
            token = await self._access_token(client)
            files = await self._list_files(client, token)
            active_file_ids: list[str] = []
            for metadata in files:
                file_id = metadata['id']
                active_file_ids.append(file_id)
                try:
                    async with client.stream(
                        'GET',
                        f'https://www.googleapis.com/drive/v3/files/{quote(file_id, safe="")}',
                        params={'alt': 'media'},
                        headers={'Authorization': f'Bearer {token}'},
                    ) as response:
                        response.raise_for_status()
                        chunks = []
                        size = 0
                        async for chunk in response.aiter_bytes():
                            size += len(chunk)
                            if size > 20 * 1024 * 1024:
                                raise DriveOrderError('Una orden supera el límite permitido de 20 MB.')
                            chunks.append(chunk)
                        content = b''.join(chunks)
                except DriveOrderError:
                    raise
                except httpx.HTTPError as exc:
                    raise DriveOrderError('No se pudo descargar una orden PDF de Google Drive.') from exc
                try:
                    order = parse_order(metadata['name'], file_id, content, suppliers, self.settings.operational_today())
                except OrderParseError as exc:
                    order = {
                        'source_id': file_id,
                        'source_name': metadata['name'],
                        'id_origen': order_origin(file_id),
                        'sha256': hashlib.sha256(content).hexdigest(),
                        'status': 'NEEDS_REVIEW',
                        'reason': str(exc),
                        'method': None,
                        'supplier': None,
                        'amount': None,
                        'cheque_date': None,
                        'cheque_type': None,
                    }
                journal.store_order(order)
            journal.retain_only_orders(active_file_ids)
        return {
            'enviada': False,
            'solo_lectura': True,
            'archivos_op_pdf_en_drive': len(files),
            'ordenes': journal.list_orders(),
        }
