from io import BytesIO
from typing import Any

from fastapi.responses import StreamingResponse
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter


def workbook_output(sheets: dict[str, tuple[list[str], list[list[Any]]]]) -> BytesIO:
    workbook = Workbook()
    workbook.remove(workbook.active)
    for name, (headers, rows) in sheets.items():
        sheet = workbook.create_sheet(name[:31] or 'Datos')
        sheet.append(headers)
        for cell in sheet[1]:
            cell.font = Font(bold=True, color='FFFFFF')
            cell.fill = PatternFill('solid', fgColor='34415E')
        for row in rows:
            safe_values = []
            for value in row:
                if isinstance(value, str) and value.startswith(('=', '+', '-', '@')):
                    value = "'" + value
                safe_values.append(value)
            sheet.append(safe_values)
        sheet.freeze_panes = 'A2'
        if headers:
            sheet.auto_filter.ref = sheet.dimensions
        for column in sheet.columns:
            values = [len(str(cell.value or '')) for cell in column]
            width = min(max(max(values, default=0) + 2, 12), 42)
            sheet.column_dimensions[get_column_letter(column[0].column)].width = width
    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    return output


def workbook_response(filename: str, sheets: dict[str, tuple[list[str], list[list[Any]]]]):
    return StreamingResponse(workbook_output(sheets),
        media_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        headers={'Content-Disposition': f'attachment; filename="{filename}"'})
