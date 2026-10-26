import asyncio
import base64
import json
import logging
import time
from datetime import date
from pathlib import Path
from urllib.parse import parse_qs
import httpx
import pytest
from argon2 import PasswordHasher
from app.auth import AccountLocked, InvalidCredentials, UserAuth
from cryptography.hazmat.primitives import hashes, serialization
from cryptography.hazmat.primitives.asymmetric import padding, rsa
from fastapi.testclient import TestClient
from pydantic import ValidationError
from app.client import BankError, CredicoopClient, signed_assertion
from app.config import Settings
from app.journal import Journal, DuplicateInstruction
from app.main import create_app
from app.models import (
    EcheqListRequest, EcheqManagementRequest, EcheqRequest, FciRedemptionRequest,
    FciSubscriptionRequest, TransferRequest, echeq_list_payload, echeq_management_payload,
    fci_instruction_payload, instruction_payload, money,
)


@pytest.fixture
def cfg(tmp_path):
    key = rsa.generate_private_key(public_exponent=65537, key_size=2048)
    path = tmp_path / 'test.pem'
    path.write_bytes(key.private_bytes(serialization.Encoding.PEM, serialization.PrivateFormat.PKCS8, serialization.NoEncryption()))
    return Settings(_env_file=None, api_key='a'*40, private_key_path=path, journal_path=tmp_path/'ops.sqlite3')


def transfer(origin='prueba-001'):
    return {'idOrigen': origin, 'cbuCuentaDebito': '1910000000000000000000',
            'beneficiarios': [{'cbuCvu': '0290023010000000555675', 'monto': '101.12',
                              'cui': '30677886370', 'nombre': 'KRAFT ELVA ROSA'}]}


def echeq():
    return {'idOrigen': 'echeq-001', 'cbuCuentaDebito': '1910000000000000000000',
            'echeqs': [{'monto': '15.79', 'fechaPago': '2026-09-02', 'motivoPago': 'PRUEBA',
                        'beneficiarioNombre': 'FUNDACION MEDICA DE MAR DEL PLATA',
                        'beneficiarioDocumento': '30546125501', 'tipoCheque': 'ECHD'}]}


def echeq_management(action='ACEPTAR', origin='echeq-management-001'):
    result = {'idOrigen': origin, 'cbuCuenta': '1910000000000000000000', 'accion': action,
              'echeqs': [{'idCheque': 'DEMOECHEQ001'}]}
    if action in ('DEPOSITAR', 'PONER-EN-CUSTODIA'):
        result['echeqs'][0].update(fechaPago='2026-09-04', monto='1500.00')
    if action == 'ENDOSAR':
        result.update(tipoEndoso='NOM', beneficiario={'documento': '30517123123', 'documentoTipo': 'CUIT'})
    return result


def fci_account():
    return {'tipoCuenta': 'ORDI', 'sucursalCuenta': '119', 'numeroCuenta': '011123/9'}


def fci_subscription(origin='fci-sub-001'):
    return {'idOrigen': origin, 'cbuCuentaDebito': '1910119655111234183598',
                        'cuentaComitente': fci_account(), 'solicitudSuscripcion': {
                            'codigoFondo': 'FCAD', 'moneda': 'ARS', 'monto': '10000.00',
                            'avanzarTestVencido': False, 'aceptarRiesgoExcedido': False}}


def fci_redemption(origin='fci-red-001'):
    return {'idOrigen': origin, 'cbuCuentaCredito': '1910119655111234183598',
                        'cuentaComitente': fci_account(), 'solicitudRescate': {
                            'codigoFondo': 'FCAD', 'moneda': 'ARS', 'cuotapartes': '500'}}


def test_jwt_signature_audience_and_unique_id(cfg):
    first, second = signed_assertion(cfg), signed_assertion(cfg)
    header, payload, signature = first.split('.')
    decode = lambda s: base64.urlsafe_b64decode(s + '='*(-len(s)%4))
    claims = json.loads(decode(payload))
    assert claims['aud'] == cfg.realm_url
    assert claims['exp'] - claims['iat'] == 60
    assert claims['iss'] == claims['sub'] == cfg.client_id
    assert first != second
    key = serialization.load_pem_private_key(cfg.private_key_path.read_bytes(), None)
    key.public_key().verify(decode(signature), (header+'.'+payload).encode(), padding.PKCS1v15(), hashes.SHA256())


def test_user_login_argon_sessions_and_admin_user_management(cfg):
    users = [{'username': 'IMRCandela', 'password_hash': PasswordHasher().hash('new-secret-passphrase'), 'role': 'admin'}]
    settings = Settings(_env_file=None, api_key='a'*40, private_key_path=cfg.private_key_path,
                        journal_path=cfg.journal_path, auth_users_json=json.dumps(users))
    bank = CredicoopClient(settings, httpx.MockTransport(lambda _: pytest.fail('No bank call expected')), interval=0)
    with TestClient(create_app(settings, bank)) as api:
        assert api.get('/auth/me').status_code == 401
        assert api.post('/login', json={'username': 'IMRCandela', 'password': 'wrong-password'}).status_code == 401
        login = api.post('/login', json={'username': 'IMRCandela', 'password': 'new-secret-passphrase'})
        assert login.status_code == 200 and 'httponly' in login.headers['set-cookie'].lower()
        assert api.get('/auth/me').json()['usuario'] == 'IMRCandela'
        assert api.get('/usuarios').status_code == 200
        created = api.post('/usuarios', json={'username': 'IMRMelisa', 'password': 'another-secure-passphrase', 'role': 'user'})
        assert created.status_code == 201
        assert api.patch('/usuarios/IMRMelisa', json={'active': False}).status_code == 200
        assert api.get('/auditoria').json()['eventos']
        assert 'password_hash' not in api.get('/usuarios').text


def test_session_requires_password_every_thirty_minutes_and_rotates_cookie(cfg):
    users = [{'username': 'operator', 'password_hash': PasswordHasher().hash('long-enough-new-password'), 'role': 'admin'}]
    settings = Settings(_env_file=None, api_key='a'*40, private_key_path=cfg.private_key_path,
                        journal_path=cfg.journal_path, auth_users_json=json.dumps(users))
    bank = CredicoopClient(settings, httpx.MockTransport(lambda _: pytest.fail('No bank call expected')), interval=0)
    app = create_app(settings, bank)
    credentials = {'username': 'operator', 'password': 'long-enough-new-password'}
    with TestClient(app) as api:
        login = api.post('/login', json=credentials)
        assert login.status_code == 200 and login.json()['revalidar_en_segundos'] == 1800
        old_token = api.cookies.get('ronanfer_session')
        with app.state.auth.connect() as conn:
            conn.execute('UPDATE web_sessions SET created_at=?,last_seen_at=?', (time.time()-1801, time.time()))
        me = api.get('/auth/me').json()
        assert me['revalidacion_requerida'] and me['revalidar_en_segundos'] == 0
        assert api.get('/usuarios').status_code == 428
        assert api.post('/transferencias', json=transfer('expired-session')).status_code == 428
        assert api.post('/login', json={**credentials, 'password': 'incorrect-password'}).status_code == 401
        assert api.get('/usuarios').status_code == 428
        assert api.post('/login', json=credentials).status_code == 200
        assert api.cookies.get('ronanfer_session') != old_token
        assert app.state.auth.session_user(old_token) is None
        assert not api.get('/auth/me').json()['revalidacion_requerida']
        assert api.get('/usuarios').status_code == 200
        with app.state.auth.connect() as conn:
            conn.execute('UPDATE web_sessions SET created_at=?', (time.time()-1801,))
        assert api.post('/logout').status_code == 200
        assert api.get('/auth/me').status_code == 401
        assert api.post('/logout').status_code == 200


def test_login_lockout_and_idle_session_expiry(tmp_path):
    password_hash = PasswordHasher().hash('long-enough-new-password')
    locked = UserAuth(tmp_path/'locked.sqlite3', json.dumps([{
        'username': 'operator', 'password_hash': password_hash, 'role': 'user'}]))
    for _ in range(4):
        with pytest.raises(InvalidCredentials):
            locked.login('operator', 'incorrect-password')
    with pytest.raises(AccountLocked):
        locked.login('operator', 'incorrect-password')
    with pytest.raises(AccountLocked):
        locked.login('operator', 'long-enough-new-password')
    sessions = UserAuth(tmp_path/'sessions.sqlite3', json.dumps([{
        'username': 'operator', 'password_hash': password_hash, 'role': 'user'}]))
    token, _ = sessions.login('operator', 'long-enough-new-password')
    with sessions.connect() as conn:
        conn.execute('UPDATE web_sessions SET last_seen_at=?', (time.time()-16*60,))
    assert sessions.session_user(token) is None


@pytest.mark.parametrize('value', ['0', '-1', 'NaN', 'Infinity', '1.001', '10000000000000', 1.2, True, '1,20'])
def test_bad_money(value):
    with pytest.raises(ValueError): money(value)


def test_money_exact():
    assert money('0.10') == '0.10'
    assert money('9999999999999.99') == '9999999999999.99'


def test_payload_adherente_dates_and_orders(cfg):
    item = TransferRequest(**transfer())
    body = instruction_payload(item, cfg.adherente, cfg.operational_today())
    assert body['numeroAdherente'] == 661395
    assert body['beneficiarios'][0]['orden'] == 0
    assert 'operadoresFirmantes' not in body
    e = instruction_payload(EcheqRequest(**echeq()), cfg.adherente, cfg.operational_today())
    assert e['echeqs'][0]['fechaPago'] == '20260902'
    assert e['echeqs'][0]['caracter'] == '1'


def test_echeq_management_and_list_payloads(cfg):
    management = echeq_management('DEPOSITAR')
    management['echeqs'][0]['cmc7'] = '19134335611012340603430010005'
    payload = echeq_management_payload(EcheqManagementRequest(**management), cfg.adherente)
    assert payload['numeroAdherente'] == cfg.adherente
    assert payload['accion'] == 'DEPOSITAR'
    assert payload['echeqs'][0]['fechaPago'] == '20260904'
    assert payload['echeqs'][0]['monto'] == '1500.00'
    assert 'operadoresFirmantes' not in payload
    listing = EcheqListRequest(idOrigen='list-001', filtro={
        'gestion': 'RECIBIDOS', 'estado': 'ACTIVO', 'fechaEmisionDesde': '2026-08-01',
        'fechaEmisionHasta': '2026-08-31', 'pagina': 2, 'limite': 50,
    })
    list_payload = echeq_list_payload(listing, cfg.adherente)
    assert list_payload['filtro']['fechaEmisionDesde'] == '20260801'
    assert list_payload['filtro']['pagina'] == 2 and list_payload['filtro']['limite'] == 50
    with pytest.raises(ValidationError):
        EcheqManagementRequest(**{**echeq_management('ACEPTAR'), 'echeqs': [{'idCheque': 'x', 'monto': '10'}]})


def test_echeq_only_allows_a_la_orden():
    data = echeq()
    data['echeqs'][0]['caracter'] = '0'
    with pytest.raises(ValidationError):
        EcheqRequest(**data)


def test_echeq_current_and_deferred_date_rules(cfg):
    data = echeq()
    data['echeqs'][0]['tipoCheque'] = 'ECHC'
    with pytest.raises(ValueError): instruction_payload(EcheqRequest(**data), cfg.adherente, cfg.operational_today())
    data['echeqs'][0]['fechaPago'] = '20260828'
    assert instruction_payload(EcheqRequest(**data), cfg.adherente, cfg.operational_today())['echeqs'][0]['fechaPago'] == '20260828'
    data['echeqs'][0]['tipoCheque'] = 'ECHD'
    with pytest.raises(ValueError): instruction_payload(EcheqRequest(**data), cfg.adherente, cfg.operational_today())


def test_transfer_date_must_be_bank_day(cfg):
    data = transfer(); data['fechaPago'] = '2026-08-29'
    with pytest.raises(ValueError): instruction_payload(TransferRequest(**data), cfg.adherente, cfg.operational_today())


def test_direct_scopes_blocked(cfg):
    with pytest.raises(ValidationError):
        Settings(_env_file=None, api_key='a'*40, scopes=cfg.scopes+' transferencias')
    with pytest.raises(ValidationError):
        Settings(_env_file=None, api_key='a'*40, scopes=cfg.scopes+' fci')


def test_fci_instruction_payload_fixes_signer_and_rescue_units(cfg):
    subscription = fci_instruction_payload(FciSubscriptionRequest(**fci_subscription()), cfg.adherente)
    assert subscription['numeroAdherente'] == 661395
    assert subscription['operadoresFirmantes'] == [{'documento': '44379155', 'documentoTipo': 'DNI'}]
    assert subscription['solicitudSuscripcion']['monto'] == '10000.00'
    redemption = fci_instruction_payload(FciRedemptionRequest(**fci_redemption()), cfg.adherente)
    assert redemption['solicitudRescate']['cuotapartes'] == 500
    assert isinstance(redemption['solicitudRescate']['cuotapartes'], int)


def test_fci_redemption_requires_exactly_one_amount_or_units():
    data = fci_redemption()
    data['solicitudRescate']['monto'] = '100.00'
    with pytest.raises(ValidationError):
        FciRedemptionRequest(**data)
    del data['solicitudRescate']['cuotapartes']
    del data['solicitudRescate']['monto']
    with pytest.raises(ValidationError):
        FciRedemptionRequest(**data)


def test_fci_movement_date_range_is_validated():
    from app.models import FciMovementsRequest
    data = {'cuentaComitente': fci_account(), 'fechaDesde': '2026-08-29', 'fechaHasta': '2026-08-28'}
    with pytest.raises(ValidationError):
        FciMovementsRequest(**data)


@pytest.mark.parametrize('linked', [[{'cbu': '1910119655111234183598'}], None])
def test_fci_accounts_logs_structure_without_sensitive_values(cfg, caplog, linked):
    payload = {'data': {'cuentasComitentes': [fci_account()], 'cuentasVinculadas': linked,
                        'perfilInversor': 'MODERADO'}}

    def handler(req):
        if req.url.path.endswith('/token'):
            assert 'fciConFirma' in parse_qs(req.content.decode())['scope'][0].split()
            return httpx.Response(200, json={'access_token': 'secret-fci-token', 'expires_in': 1800})
        return httpx.Response(200, json=payload)

    async def run():
        client = CredicoopClient(cfg, httpx.MockTransport(handler), interval=0)
        try:
            assert await client.fci_accounts() == payload
        finally:
            await client.close()

    with caplog.at_level(logging.INFO, logger='app.client'):
        asyncio.run(run())
    messages = '\n'.join(record.getMessage() for record in caplog.records if record.name == 'app.client')
    assert 'fciConFirma configurado=True' in messages
    assert 'RESPONSE FCI cuentas-comitentes: estructura=' in messages
    assert 'cuentasComitentes' in messages and 'cuentasVinculadas' in messages
    assert 'secret-fci-token' not in messages
    assert '1910119655111234183598' not in messages
    assert fci_account()['numeroCuenta'] not in messages
    assert 'MODERADO' not in messages
    if linked is None:
        assert 'null' in messages and 'formato diferente al documentado' in messages


def test_fci_routes_use_only_con_firma_and_documented_paths(cfg):
    seen = []

    def handler(req):
        seen.append(req)
        if req.url.path.endswith('/token'):
            form = parse_qs(req.content.decode())
            assert 'fciConFirma' in form['scope'][0].split()
            assert 'fci' not in form['scope'][0].split()
            return httpx.Response(200, json={'access_token': 'test-token', 'expires_in': 1800})
        if req.url.path.endswith('/cuentas-comitentes'):
            assert req.url.params['numeroAdherente'] == '661395'
            assert req.url.params['idOrigen']
            return httpx.Response(200, json={'data': {'cuentasComitentes': [fci_account()]}})
        if req.url.path.endswith('/fondos'):
            return httpx.Response(200, json={'data': {'detalleFondo': [{'codigo': 'FCAD'}]}})
        if req.url.path.endswith('/cuenta-comitente-saldos'):
            assert json.loads(req.content)['cuentaComitente'] == fci_account()
            return httpx.Response(200, json={'data': {'saldoDetalle': []}})
        if req.url.path.endswith('/cuenta-comitente-movimientos'):
            body = json.loads(req.content)
            assert body['fechaDesde'] == '20260801' and body['fechaHasta'] == '20260828'
            return httpx.Response(200, json={'data': {'movimientos': []}})
        if req.url.path.endswith('/cuenta-comitente-movimiento-detalle'):
            assert req.url.params['formula'] == '7018151'
            return httpx.Response(200, json={'data': {'detalle': []}})
        if req.method == 'POST':
            assert req.url.path in {
                '/api/fci/v1/ConFirma/suscripcion', '/api/fci/v1/ConFirma/rescate',
            }
            body = json.loads(req.content)
            assert body['operadoresFirmantes'] == [{'documento': '44379155', 'documentoTipo': 'DNI'}]
            return httpx.Response(200, json={'data': {'idOperacion': 77, 'estadoOperacion': {'descripcion': 'Enviada a la firma'}}})
        if req.url.path.endswith(('/suscripcion', '/rescate')):
            assert req.url.params.get('idOrigen') or req.url.params.get('idOperacion')
            return httpx.Response(200, json={'data': {'idOperacion': 77}})
        return httpx.Response(404, json={'error': 'Ruta de prueba no esperada'})

    b = CredicoopClient(cfg, httpx.MockTransport(handler), interval=0)
    with TestClient(create_app(cfg, b)) as api:
        headers = {'X-API-Key': 'a'*40}
        assert api.get('/fci/cuentas-comitentes', headers=headers).status_code == 200
        assert api.get('/fci/fondos', headers=headers).status_code == 200
        assert api.post('/fci/saldos', headers=headers, json={'cuentaComitente': fci_account()}).status_code == 200
        movements = {'cuentaComitente': fci_account(), 'fechaDesde': '2026-08-01', 'fechaHasta': '2026-08-28'}
        assert api.post('/fci/movimientos', headers=headers, json=movements).status_code == 200
        assert api.get('/fci/movimientos/detalle?tipo=RESC&numero=500&sucursal=0123&formula=7018151',
                       headers=headers).status_code == 200
        subscription = fci_subscription()
        preview = api.post('/fci/suscripciones/previsualizar', headers=headers, json=subscription)
        assert preview.status_code == 200 and preview.json()['enviada'] is False
        assert preview.json()['payload_banco']['operadoresFirmantes'][0]['documento'] == '44379155'
        sent = api.post('/fci/suscripciones', headers=headers, json=subscription)
        assert sent.status_code == 200 and sent.json()['modalidad'] == 'ConFirma'
        assert api.post('/fci/suscripciones', headers=headers, json=subscription).status_code == 409
        assert api.post('/fci/rescates', headers=headers, json=fci_redemption()).status_code == 200
        assert api.get('/fci/suscripciones/estado?id_operacion=77', headers=headers).status_code == 200
        assert api.get('/fci/rescates/estado?id_origen=fci-red-001', headers=headers).status_code == 200
        assert api.get('/operaciones/fci_suscripcion/fci-sub-001', headers=headers).status_code == 200
    assert len([req for req in seen if req.method == 'POST' and '/ConFirma/' in req.url.path]) == 2


def test_fci_routes_explain_missing_scope_without_blocking_existing_service(cfg):
    scopes = cfg.scopes.replace(' fciConFirma', '')
    settings = Settings(_env_file=None, api_key='a'*40, scopes=scopes, private_key_path=cfg.private_key_path,
                       journal_path=cfg.journal_path)
    bank = CredicoopClient(settings, httpx.MockTransport(lambda _: pytest.fail('Banco no debe ser contactado')), interval=0)
    with TestClient(create_app(settings, bank)) as api:
        headers = {'X-API-Key': 'a'*40}
        assert api.get('/salud').status_code == 200
        response = api.get('/fci/cuentas-comitentes', headers=headers)
        assert response.status_code == 503 and 'fciConFirma' in response.json()['detail']


def test_fci_api_business_error_is_not_returned_as_success(cfg):
    def handler(req):
        if req.url.path.endswith('/token'):
            return httpx.Response(200, json={'access_token': 't', 'expires_in': 1800})
        return httpx.Response(200, json={'error': {'codigo': 'APIE-1101', 'descripcion': 'Parametros invalidos'}})

    async def go():
        bank = CredicoopClient(cfg, httpx.MockTransport(handler), interval=0)
        with pytest.raises(BankError, match='informó errores'):
            await bank.fci_funds()
        await bank.close()
    asyncio.run(go())


def test_local_key_required():
    with pytest.raises(ValidationError): Settings(_env_file=None, api_key='')


def test_production_date_can_be_unset(monkeypatch):
    monkeypatch.setenv('CREDICOOP_FECHA_OPERATIVA', 'null')
    cfg = Settings(_env_file=None, api_key='a'*40)
    assert cfg.fecha_operativa is None


def test_batch_validation():
    data = transfer()
    data['beneficiarios'].append({**data['beneficiarios'][0], 'cbuCvu': '0000000000000000000000'})
    with pytest.raises(ValidationError): TransferRequest(**data)
    data = echeq()
    data['echeqs'].append({**data['echeqs'][0], 'numeroCheque': 1})
    with pytest.raises(ValidationError): EcheqRequest(**data)


def test_journal_survives_restart_and_changed_body(cfg):
    first = Journal(cfg.journal_path); body = transfer()
    first.reserve('transferencia', body)
    second = Journal(cfg.journal_path)
    body['beneficiarios'][0]['monto'] = '20.00'
    with pytest.raises(DuplicateInstruction): second.reserve('transferencia', body)
    assert second.get('transferencia', 'prueba-001')['estado'] == 'ENVIO_INICIADO'
    second.record_event('candela', 'login', '/login', 200)
    from openpyxl import load_workbook
    report = next(cfg.journal_path.parent.glob('auditoria-*.xlsx'))
    workbook = load_workbook(report, read_only=True, data_only=True)
    assert {'transferencia', 'actividad'} <= set(workbook.sheetnames)
    assert workbook['transferencia'].max_row == 2 and workbook['actividad'].max_row == 2
    workbook.close()


def test_operations_filter_summary_and_saved_bank_status(cfg):
    def handler(req):
        if req.url.path.endswith('/token'):
            return httpx.Response(200, json={'access_token': 't', 'expires_in': 1800})
        if req.method == 'POST':
            return httpx.Response(200, json={'data': {'idOperacion': '33328471',
                                                       'estadoOperacion': {'descripcion': 'PENDIENTE'}}})
        assert req.url.params['idOrigen'] == 'prueba-status-001'
        return httpx.Response(200, json={'data': {'idOperacion': '33328471',
                                                   'estadoOperacion': {'descripcion': 'FIRMADA'}}})

    with TestClient(create_app(cfg, CredicoopClient(cfg, httpx.MockTransport(handler), interval=0))) as api:
        headers = {'X-API-Key': 'a'*40}
        sent = api.post('/transferencias', headers=headers, json=transfer('prueba-status-001'))
        assert sent.status_code == 200
        operations = api.get('/operaciones?tipo=transferencia', headers=headers).json()['operaciones']
        assert len(operations) == 1
        assert operations[0]['resumen']['beneficiario'] == 'KRAFT ELVA ROSA'
        assert operations[0]['resumen']['importe'] == '101.12'
        status = api.get('/transferencias/estado?id_origen=prueba-status-001', headers=headers)
        assert status.status_code == 200
        saved = api.get('/operaciones?tipo=transferencia', headers=headers).json()['operaciones'][0]
        assert saved['resultado']['consulta_estado']['data']['estadoOperacion']['descripcion'] == 'FIRMADA'
        exported = api.get('/operaciones/export.xlsx', headers=headers)
        assert exported.status_code == 200
        assert 'spreadsheetml' in exported.headers['content-type']


def test_api_only_con_firma_and_no_double_send(cfg):
    seen = []
    def handler(req):
        seen.append(req)
        if req.url.path.endswith('/token'):
            form = parse_qs(req.content.decode())
            assert 'transferenciasConFirma' in form['scope'][0]
            assert 'transferencias' not in form['scope'][0].split()
            return httpx.Response(200, json={'access_token': 'test-token', 'expires_in': 1800})
        if req.method == 'POST':
            assert '/ConFirma/' in req.url.path
            data = json.loads(req.content)
            assert data['numeroAdherente'] == 661395
            return httpx.Response(200, json={'data': {'idOperacion': 77, 'estadoOperacion': {'descripcion': 'Enviada a la firma'}}})
        return httpx.Response(200, json={'clarifCuentas': [{'nroCuenta': '00123', 'saldo': '10.20', 'CBU': '1910000000000000000000'}]})
    b = CredicoopClient(cfg, httpx.MockTransport(handler), interval=0)
    with TestClient(create_app(cfg, b)) as api:
        assert api.get('/cuentas').status_code == 401
        headers = {'X-API-Key': 'a'*40}
        assert api.get('/cuentas/00123/saldo', headers=headers).json()['saldo'] == '10.20'
        count = len(seen)
        assert api.post('/transferencias/previsualizar', headers=headers, json=transfer()).json()['enviada'] is False
        assert len(seen) == count
        r = api.post('/transferencias', headers=headers, json=transfer())
        assert r.status_code == 200
        assert r.json()['respuesta_banco']['data']['estadoOperacion']['descripcion'] == 'Enviada a la firma'
        assert api.post('/transferencias', headers=headers, json=transfer()).status_code == 409
        assert api.post('/echeqs', headers=headers, json=echeq()).status_code == 200
    assert len([x for x in seen if x.url.path.endswith('/token')]) == 1
    assert len([x for x in seen if '/ConFirma/' in x.url.path]) == 2


def test_echeq_list_and_manage_use_con_firma_and_are_deduplicated(cfg):
    seen = []
    def handler(req):
        seen.append(req)
        if req.url.path.endswith('/token'):
            form = parse_qs(req.content.decode())
            assert 'echeqConFirma' in form['scope'][0].split()
            assert 'echeq' not in form['scope'][0].split()
            return httpx.Response(200, json={'access_token': 'test-token', 'expires_in': 1800})
        if req.url.path.endswith('/lista-cheques'):
            body = json.loads(req.content)
            assert body['numeroAdherente'] == cfg.adherente
            assert body['filtro']['fechaEmisionDesde'] == '20260801'
            return httpx.Response(200, json={'data': {'echeqs': [{'chequeId': 'DEMOECHEQ001', 'estado': 'ACTIVO'}], 'totalCheques': 1}})
        if req.url.path.endswith('/ConFirma/gestion'):
            body = json.loads(req.content)
            assert body['accion'] == 'ACEPTAR' and body['numeroAdherente'] == cfg.adherente
            return httpx.Response(200, json={'data': {'idOrigen': body['idOrigen'],
                'estadoOperacion': {'descripcion': 'En procesamiento'}}})
        return httpx.Response(404, json={'error': 'Ruta de prueba no esperada'})

    bank = CredicoopClient(cfg, httpx.MockTransport(handler), interval=0)
    with TestClient(create_app(cfg, bank)) as api:
        headers = {'X-API-Key': 'a'*40}
        before = len(seen)
        preview_body = echeq_management('ACEPTAR')
        preview = api.post('/echeqs/gestionar/previsualizar', headers=headers, json=preview_body)
        assert preview.status_code == 200 and preview.json()['enviada'] is False and len(seen) == before
        listing = api.post('/echeqs/listar', headers=headers, json={'idOrigen': 'list-001', 'filtro': {
            'fechaEmisionDesde': '2026-08-01', 'fechaEmisionHasta': '2026-08-31', 'pagina': 1, 'limite': 20}})
        assert listing.status_code == 200 and listing.json()['data']['totalCheques'] == 1
        sent = api.post('/echeqs/gestionar', headers=headers, json=preview_body)
        assert sent.status_code == 200
        assert sent.json()['respuesta_banco']['data']['estadoOperacion']['descripcion'] == 'En procesamiento'
        assert api.post('/echeqs/gestionar', headers=headers, json=preview_body).status_code == 409
        assert api.get('/operaciones/echeq_gestion/echeq-management-001', headers=headers).status_code == 200
    assert any(req.url.path.endswith('/api/echeq/v1/ConFirma/gestion') for req in seen)


def test_timeout_blocks_retry(cfg):
    count = 0
    def handler(req):
        nonlocal count
        if req.url.path.endswith('/token'): return httpx.Response(200, json={'access_token': 't', 'expires_in': 30})
        count += 1
        raise httpx.ReadTimeout('test', request=req)
    b = CredicoopClient(cfg, httpx.MockTransport(handler), interval=0)
    with TestClient(create_app(cfg, b)) as api:
        headers = {'X-API-Key': 'a'*40}
        assert api.post('/echeqs', headers=headers, json=echeq()).status_code == 504
        assert api.get('/operaciones/echeq/echeq-001', headers=headers).json()['estado'] == 'RESULTADO_INCIERTO'
        assert api.post('/echeqs', headers=headers, json=echeq()).status_code == 409
    assert count == 1


def test_post_401_is_not_retried(cfg):
    count = 0
    def handler(req):
        nonlocal count
        if req.url.path.endswith('/token'): return httpx.Response(200, json={'access_token': 't', 'expires_in': 30})
        count += 1
        return httpx.Response(401, json={'codigo': 'APIE-0010'})
    async def go():
        b = CredicoopClient(cfg, httpx.MockTransport(handler), interval=0)
        with pytest.raises(BankError): await b.request('POST', '/api/echeq/v1/ConFirma/emision', body={})
        await b.close()
    asyncio.run(go()); assert count == 1


def test_get_401_refreshes_once(cfg):
    tokens = 0
    def handler(req):
        nonlocal tokens
        if req.url.path.endswith('/token'):
            tokens += 1
            return httpx.Response(200, json={'access_token': f't{tokens}', 'expires_in': 1800})
        if req.headers['Authorization'] == 'Bearer t1': return httpx.Response(401, json={})
        return httpx.Response(200, json={'clarifCuentas': []})
    async def go():
        b = CredicoopClient(cfg, httpx.MockTransport(handler), interval=0)
        assert await b.accounts() == {'clarifCuentas': []}
        await b.close()
    asyncio.run(go()); assert tokens == 2


def test_movement_splitting_headers_and_partial_day(cfg):
    calls = []
    def handler(req):
        if req.url.path.endswith('/token'): return httpx.Response(200, json={'access_token': 't', 'expires_in': 1800})
        a, z = req.url.params['fechaDesde'], req.url.params['fechaHasta']; calls.append((a,z))
        if a != z: return httpx.Response(200, json={'alerta': [{'mensaje': 'parcial'}], 'consMovCtas': [{'descripcion': 'NO CONSERVAR'}]})
        body = {'consMovCtas': [{'descripcion': 'ENCABEZADO', 'saldo': 42}, {'descripcion': a, 'idTransaccion': a}]}
        if a == '20260827': body['alerta'] = [{'mensaje': 'parcial por dia'}]
        return httpx.Response(200, json=body)
    async def go():
        b = CredicoopClient(cfg, httpx.MockTransport(handler), interval=0)
        r = await b.movements('00123', date(2026,8,27), date(2026,8,28))
        assert len(r['movimientos']) == 2 and len(r['encabezados']) == 2
        assert not r['completa'] and len(r['rangos_incompletos']) == 1
        assert [x['descripcion'] for x in r['movimientos']] == ['20260828','20260827']
        await b.close()
    asyncio.run(go()); assert len(calls) == 3


def test_movements_default_current_argentina(cfg, monkeypatch):
    monkeypatch.setattr('app.main.argentina_today', lambda: date(2026,10,5))
    calls = []
    def handler(req):
        if req.url.path.endswith('/token'): return httpx.Response(200, json={'access_token': 't', 'expires_in': 1800})
        calls.append(req.url.params); return httpx.Response(200, json={'consMovCtas': []})
    b = CredicoopClient(cfg, httpx.MockTransport(handler), interval=0)
    with TestClient(create_app(cfg,b)) as api:
        h = {'X-API-Key': 'a'*40}
        r = api.get('/cuentas/00123/movimientos?fecha_desde=2026-08-01', headers=h)
        assert r.status_code == 200 and r.json()['fecha_hasta'] == '2026-10-05'
        assert calls[0]['fechaHasta'] == '20261005'
        assert api.get('/cuentas/00123/movimientos?fecha_desde=2026-10-06', headers=h).status_code == 422


def test_bad_response_never_claims_sent(cfg):
    def handler(req):
        if req.url.path.endswith('/token'): return httpx.Response(200, json={'access_token': 't', 'expires_in': 1800})
        return httpx.Response(200, json={})
    with TestClient(create_app(cfg, CredicoopClient(cfg, httpx.MockTransport(handler), interval=0))) as api:
        h = {'X-API-Key': 'a'*40}
        r = api.post('/transferencias', headers=h, json=transfer())
        assert r.status_code == 502 and r.json()['resultado_incierto']
        assert api.post('/transferencias', headers=h, json=transfer()).status_code == 409


def test_movements_request_and_response_are_logged(cfg, caplog):
    def handler(req):
        if req.url.path.endswith('/token'): return httpx.Response(200, json={'access_token': 't', 'expires_in': 1800})
        if req.url.path.endswith('/listaCuentas'):
            assert req.url.path == '/api/cuentas/v2/listaCuentas'
            return httpx.Response(200, json={'clarifCuentas': []})
        return httpx.Response(200, json={'consMovCtas': [{'descripcion': 'ENCABEZADO', 'saldo': 42}]})
    async def go():
        b = CredicoopClient(cfg, httpx.MockTransport(handler), interval=0)
        await b.movements('00123', date(2026,8,28), date(2026,8,28))
        await b.accounts()
        await b.close()
    with caplog.at_level(logging.INFO, logger='app.client'):
        asyncio.run(go())
    requests = [r.getMessage() for r in caplog.records if 'REQUEST movimientos' in r.getMessage()]
    responses = [r.getMessage() for r in caplog.records if 'RESPONSE movimientos' in r.getMessage()]
    # Sólo la consulta de movimientos se loguea; /listaCuentas queda fuera.
    assert len(requests) == 1 and len(responses) == 1
    assert 'GET /api/cuentas/v2/00123/movimientos' in requests[0]
    assert '20260828' in requests[0] and 'body=null' in requests[0]
    assert 'status=200' in responses[0] and 'consMovCtas' in responses[0]


def test_movements_logging_includes_fci_body_and_can_be_disabled(cfg, caplog):
    payload = {'cuentaComitente': fci_account(), 'fechaDesde': '20260801', 'fechaHasta': '20260828'}
    def handler(req):
        if req.url.path.endswith('/token'): return httpx.Response(200, json={'access_token': 't', 'expires_in': 1800})
        return httpx.Response(200, json={'data': {'movimientos': [], 'totalMovimientos': 0}})
    async def run(client):
        return await client.request('POST', '/api/fci/v1/cuenta-comitente-movimientos', body=payload)
    with caplog.at_level(logging.INFO, logger='app.client'):
        b = CredicoopClient(cfg, httpx.MockTransport(handler), interval=0)
        asyncio.run(run(b)); asyncio.run(b.close())
    logged = [r.getMessage() for r in caplog.records if 'REQUEST movimientos' in r.getMessage()]
    assert len(logged) == 1 and 'POST /api/fci/v1/cuenta-comitente-movimientos' in logged[0]
    body = logged[0].split('body=', 1)[1]
    assert '"fechaDesde": "20260801"' in body and '"numeroCuenta": "011123/9"' in body
    cfg.log_movimientos = False
    caplog.clear()
    with caplog.at_level(logging.INFO, logger='app.client'):
        b = CredicoopClient(cfg, httpx.MockTransport(handler), interval=0)
        asyncio.run(run(b)); asyncio.run(b.close())
    assert not any('movimientos [' in r.getMessage() for r in caplog.records)
