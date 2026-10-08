"""Ejecuta consultas reales de homologación; nunca envía pagos ni gestiones."""

import argparse
import csv
import json
import re
import sys
from datetime import date, timedelta
from pathlib import Path
from uuid import NAMESPACE_URL, uuid4, uuid5

import httpx
from openpyxl import load_workbook
from fastapi.testclient import TestClient

PROJECT_ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(PROJECT_ROOT))

from app.client import CredicoopClient
from app.config import Settings, argentina_today
from app.main import create_app


def local_request(api, recorder, method, path, *, params=None, body=None, function=None):
    headers = {'X-API-Key': api.app.state.bank.settings.api_key.get_secret_value()}
    try:
        response = api.request(method, path, params=params, json=body, headers=headers)
    except httpx.HTTPError as exc:
        request = httpx.Request(
            method, httpx.URL('http://testserver').join(path),
            params=params, json=body, headers=headers,
        )
        recorder.record(
            function or f'FastAPI {method} {path}', 'X-API-Key local',
            request, error=f'{type(exc).__name__}: {exc}',
        )
        return None
    recorder.record(function or f'FastAPI {method} {path}', 'X-API-Key local', response.request, response)
    return response


def load_test_beneficiaries(path: Path):
    with path.open(encoding='utf-8-sig', newline='') as stream:
        return list(csv.DictReader(stream, delimiter=';'))


def bank_data(response):
    try:
        body = response.json()
    except (AttributeError, ValueError):
        return None
    return body.get('data') if isinstance(body, dict) else None


def beneficiary_matches(response, beneficiary, kind):
    if response is None or response.status_code != 200:
        return False
    data = bank_data(response)
    field = 'cbuCvu' if kind == 'transferencia' else 'documento'
    expected = beneficiary.get(field, '')

    def contains_match(value):
        if isinstance(value, dict):
            return (str(value.get(field, '')) == expected
                    or any(contains_match(child) for child in value.values()))
        if isinstance(value, list):
            return any(contains_match(child) for child in value)
        return False

    return bool(expected) and contains_match(data)


def business_error_code(response):
    body = response_json(response)
    if not isinstance(body, dict):
        return None
    pending = [body.get('detalle_banco')]
    while pending:
        item = pending.pop()
        if isinstance(item, dict):
            code = item.get('codigo')
            if code is not None:
                return str(code)
            pending.extend(item.values())
        elif isinstance(item, list):
            pending.extend(item)
    return None


def transfer_master_cbus(path: Path):
    if not path.is_file():
        return set()
    workbook = load_workbook(path, read_only=True, data_only=True)
    values = set()
    try:
        for sheet in workbook.worksheets:
            for row in sheet.iter_rows(min_row=2, values_only=True):
                if row and row[0] is not None:
                    cbu = re.sub(r'\D', '', str(row[0]))
                    if len(cbu) == 22:
                        values.add(cbu)
    finally:
        workbook.close()
    return values


def response_json(response):
    if response is None:
        return None
    try:
        return response.json()
    except ValueError:
        return {'respuesta_no_json': response.text}


def account_rows(response):
    data = response_json(response)
    if not isinstance(data, dict):
        return []
    accounts = data.get('clarifCuentas', data.get('clarifcuentas', []))
    return accounts if isinstance(accounts, list) else []


def cheque_rows(response):
    data = bank_data(response)
    rows = data.get('echeqs', []) if isinstance(data, dict) else []
    return rows if isinstance(rows, list) else []


def write_preview(recorder, payloads, notes):
    if recorder.directory is None:
        raise RuntimeError('No se creó el directorio de evidencias.')
    (recorder.directory / 'PREVIAS_OPERACIONES.json').write_text(
        json.dumps({'estado': 'NO ENVIADAS', 'operaciones': payloads, 'notas': notes},
                   ensure_ascii=False, indent=2),
        encoding='utf-8',
    )


def run():
    parser = argparse.ArgumentParser(description='Consultas reales y vistas previas, sin envíos de pagos/gestión.')
    parser.add_argument('--movimientos-desde', type=date.fromisoformat)
    parser.add_argument('--movimientos-hasta', type=date.fromisoformat)
    args = parser.parse_args()
    if (args.movimientos_desde is None) != (args.movimientos_hasta is None):
        parser.error('Indicar juntas --movimientos-desde y --movimientos-hasta.')

    cfg = Settings()
    if cfg.base_url != 'https://homoapibccl.bancocredicoop.coop' or not cfg.is_homologation:
        raise RuntimeError('Ejecución cancelada: el destino no es el servidor exacto de homologación.')
    if not cfg.api_key.get_secret_value() or not cfg.private_key_path.is_file():
        raise RuntimeError('Faltan las credenciales locales de homologación.')

    bank = CredicoopClient(cfg)
    recorder = bank.evidence
    if recorder is None:
        raise RuntimeError('La captura HTTP de evidencias no está habilitada.')
    app = create_app(cfg, bank)
    results = []
    found_beneficiaries = {}
    lists = {}
    chosen_accounts = []
    now = argentina_today()
    movement_start = args.movimientos_desde or now - timedelta(days=30)
    movement_end = args.movimientos_hasta or now

    with TestClient(app) as api:
        configuration = local_request(
            api, recorder, 'GET', '/configuracion', function='Consultar configuración local',
        )
        results.append(('configuracion', configuration.status_code if configuration else None))
        accounts_response = local_request(
            api, recorder, 'GET', '/cuentas', function='Consultar cuentas habilitadas',
        )
        results.append(('cuentas', accounts_response.status_code if accounts_response else None))
        chosen_accounts = account_rows(accounts_response) if accounts_response else []

        for account in chosen_accounts:
            number = str(account.get('nroCuenta', ''))
            if not number.isdigit():
                recorder.note(
                    'Consultar movimientos',
                    'PRUEBA PENDIENTE',
                    'La respuesta de cuentas no incluyó un nroCuenta numérico; no se usó el CBU.',
                )
                continue
            movements = local_request(
                api, recorder, 'GET', f'/cuentas/{number}/movimientos',
                params={
                    'fecha_desde': movement_start.isoformat(),
                    'fecha_hasta': movement_end.isoformat(),
                },
                function=f'Consultar movimientos nroCuenta {number}',
            )
            results.append((f'movimientos:{number}', movements.status_code if movements else None))

        sample_path = PROJECT_ROOT / 'ejemplos' / 'beneficiarios_homologacion.csv'
        beneficiaries = load_test_beneficiaries(sample_path) if sample_path.is_file() else []
        transfer_master = PROJECT_ROOT / 'Beneficiarios_Credicoop.xlsx'
        bank_workbook = PROJECT_ROOT / 'BeneficiariosHomologacionAPI EMPRESA(3).xls'
        if not bank_workbook.exists():
            recorder.note(
                'Maestro bancario de homologación',
                'PRUEBA PENDIENTE',
                'No se encontró BeneficiariosHomologacionAPI EMPRESA(3).xls. Se usarán para las consultas '
                'los registros disponibles en ejemplos/beneficiarios_homologacion.csv; no se asumirá que '
                'una ausencia en Beneficiarios_Credicoop.xlsx implica ausencia en la agenda bancaria.',
            )
        if not transfer_master.exists():
            recorder.note(
                'Maestro local de beneficiarios',
                'PRUEBA PENDIENTE',
                'No se encontró Beneficiarios_Credicoop.xlsx para comparar el maestro local.',
            )
        transfer_master_values = transfer_master_cbus(transfer_master)

        for item in beneficiaries:
            kind = item.get('tipo', '')
            key = item.get('cbuCvu', '') if kind == 'transferencia' else item.get('documento', '')
            if kind == 'transferencia' and key:
                response = local_request(
                    api, recorder, 'GET', '/beneficiarios/transferencias',
                    params={'cbu_cvu': key},
                    function=f'Consultar beneficiario de transferencia {key}',
                )
            elif kind == 'echeq' and key:
                response = local_request(
                    api, recorder, 'GET', '/beneficiarios/echeqs',
                    params={
                        'documento': key,
                        'documento_tipo': item.get('documentoTipo', 'CUIT').split('/')[0],
                    },
                    function=f'Consultar beneficiario de eCheq {key}',
                )
            else:
                continue
            exists = beneficiary_matches(response, item, kind)
            found_beneficiaries[(kind, key)] = exists
            result = 'ALTA NO NECESARIA' if exists else (
                'PRUEBA PENDIENTE' if response is None or response.status_code == 200
                else 'ERROR BANCARIO'
            )
            master_note = (
                'Está en el maestro local de transferencia.' if kind == 'transferencia'
                and key in transfer_master_values else
                'No aparece en el maestro local de transferencia; esto no demuestra que falte en la agenda bancaria.'
                if kind == 'transferencia' else
                'La planilla local disponible sólo contiene el maestro de transferencias.'
            )
            if not exists and item.get('uso') == 'alta' and business_error_code(response) in {
                'APIE-7017', 'APIE-8014',
            }:
                origin = str(uuid5(NAMESPACE_URL, f'credicoop-homologacion-beneficiario:{kind}:{key}'))
                registration = {'idOrigen': origin, 'beneficiarios': []}
                if kind == 'transferencia':
                    registration['beneficiarios'].append({
                        'orden': 0,
                        'cbuCvu': item['cbuCvu'],
                        'moneda': 'ARS',
                        'cui': item['documento'],
                        'nombre': item['nombre'],
                        'visibleBI': True,
                    })
                    registration_path = '/beneficiarios/transferencias'
                else:
                    registration['beneficiarios'].append({
                        'orden': 0,
                        'documento': item['documento'],
                        'documentoTipo': item['documentoTipo'],
                    })
                    registration_path = '/beneficiarios/echeqs'
                added = local_request(
                    api, recorder, 'POST', registration_path, body=registration,
                    function=f'Alta condicional beneficiario {kind} {key}',
                )
                results.append((f'alta:{kind}:{key}', added.status_code if added else None))
                response = local_request(
                    api, recorder, 'GET',
                    '/beneficiarios/transferencias' if kind == 'transferencia' else '/beneficiarios/echeqs',
                    params=(
                        {'cbu_cvu': key} if kind == 'transferencia' else
                        {'documento': key, 'documento_tipo': item.get('documentoTipo', 'CUIT')}
                    ),
                    function=f'Verificar alta beneficiario {kind} {key}',
                )
                exists = beneficiary_matches(response, item, kind)
                found_beneficiaries[(kind, key)] = exists
                result = 'ALTA CONFIRMADA' if added and added.status_code == 200 and exists else (
                    'RESULTADO INCIERTO' if added and added.status_code == 504 else
                    'PRUEBA PENDIENTE'
                )
            recorder.note(
                f'Resultado de agenda {kind} {key}',
                result,
                ('La consulta devolvió el beneficiario esperado. Alta no necesaria. ' + master_note)
                if exists and result == 'ALTA NO NECESARIA' else
                ('El banco devolvió explícitamente el código de beneficiario ausente documentado; se hizo una sola '
                 'alta con idOrigen estable y se volvió a consultar. ' + master_note)
                if result == 'ALTA CONFIRMADA' else
                ('La consulta no confirmó de forma inequívoca la existencia o ausencia. No se hizo el alta. '
                 + master_note),
            )
            results.append((f'beneficiario:{kind}:{key}', response.status_code if response else None))

        for management_type in ('GENERADOS', 'RECIBIDOS'):
            payload = {'idOrigen': str(uuid4()), 'filtro': {'gestion': management_type}}
            response = local_request(
                api, recorder, 'POST', '/echeqs/lista', body=payload,
                function=f'Listar eCheqs {management_type}',
            )
            lists[management_type] = response
            results.append((f'lista_echeqs:{management_type}', response.status_code if response else None))

        verified_account = next(
            (a for a in chosen_accounts
             if isinstance(a.get('CBU', a.get('cbu')), str)
             and len(a.get('CBU', a.get('cbu'))) == 22
             and a.get('CBU', a.get('cbu')).isdigit()),
            None,
        )
        test_transfer = next(
            (b for b in beneficiaries if b.get('tipo') == 'transferencia'
             and b.get('uso') == 'emision'
             and found_beneficiaries.get(('transferencia', b.get('cbuCvu')))),
            None,
        )
        test_echeq = next(
            (b for b in beneficiaries if b.get('tipo') == 'echeq'
             and b.get('uso') == 'emision'
             and found_beneficiaries.get(('echeq', b.get('documento')))),
            None,
        )
        payloads = []
        notes = [
            'No se envió ninguna transferencia, emisión ni gestión.',
            'La fecha operativa del banco no se puede inferir sólo de la fecha local. Confirmarla con el banco antes de usar los payloads.',
            'No firmar ni activar operaciones en BIE durante esta ejecución.',
        ]
        if verified_account and test_transfer:
            preview = {
                'idOrigen': str(uuid4()),
                'cbuCuentaDebito': verified_account.get('CBU', verified_account.get('cbu')),
                'beneficiarios': [{
                    'cbuCvu': test_transfer['cbuCvu'],
                    'esCuentaPropia': 'N',
                    'monto': '1.00',
                    'moneda': 'ARS',
                    'concepto': 'VAR',
                    'cui': test_transfer['documento'],
                    'nombre': test_transfer['nombre'],
                }],
            }
            response = local_request(
                api, recorder, 'POST', '/transferencias/previsualizar', body=preview,
                function='Previsualizar transferencia ARS 1.00 - NO ENVIADA',
            )
            results.append(('preview_transferencia', response.status_code if response else None))
            payloads.append({
                'funcion': 'Transferir con firma',
                'resultado': 'PENDIENTE_CONFIRMACION_CONJUNTA',
                'payload_banco': (response_json(response) or {}).get('payload_banco') if response else None,
                'beneficiario': test_transfer,
                'importe': 'ARS 1.00',
                'idOrigen': preview['idOrigen'],
            })
        else:
            recorder.note(
                'Preparar transferencia ARS 1.00',
                'PRUEBA PENDIENTE',
                'No hay simultáneamente una cuenta con CBU devuelta por el banco y un beneficiario de emisión '
                'confirmado por la consulta bancaria.',
            )

        if verified_account and test_echeq:
            preview = {
                'idOrigen': str(uuid4()),
                'cbuCuentaDebito': verified_account.get('CBU', verified_account.get('cbu')),
                'echeqs': [{
                    'monto': '1.00',
                    'fechaPago': (now + timedelta(days=1)).isoformat(),
                    'motivoPago': 'PRUEBA HOMOLOGACION',
                    'caracter': '1',
                    'modo': '1',
                    'beneficiarioNombre': test_echeq['nombre'],
                    'beneficiarioDocumentoTipo': test_echeq['documentoTipo'],
                    'beneficiarioDocumento': test_echeq['documento'],
                    'concepto': 'VAR',
                    'tipoCheque': 'ECHD',
                }],
            }
            response = local_request(
                api, recorder, 'POST', '/echeqs/previsualizar', body=preview,
                function='Previsualizar emisión eCheq ARS 1.00 - NO ENVIADA',
            )
            results.append(('preview_echeq', response.status_code if response else None))
            payloads.append({
                'funcion': 'Emitir eCheq con firma',
                'resultado': 'PENDIENTE_CONFIRMACION_CONJUNTA',
                'payload_banco': (response_json(response) or {}).get('payload_banco') if response else None,
                'beneficiario': test_echeq,
                'importe': 'ARS 1.00',
                'fechaPago_propuesta_no_verificada': (now + timedelta(days=1)).isoformat(),
                'idOrigen': preview['idOrigen'],
            })
        else:
            recorder.note(
                'Preparar emisión eCheq ARS 1.00',
                'PRUEBA PENDIENTE',
                'No hay simultáneamente una cuenta con CBU devuelta por el banco y un beneficiario eCheq '
                'de emisión confirmado por la consulta bancaria.',
            )

        eligible = []
        for row in cheque_rows(lists.get('RECIBIDOS')):
            state = row.get('estado')
            cheque_id = row.get('chequeId')
            cmc7 = row.get('cmc7completo')
            if state == 'EMITIDO-PENDIENTE' and (cheque_id or cmc7):
                eligible.append({'accion': 'ACEPTAR', 'cheque': row})
            elif (state == 'ACTIVO' and (cheque_id or cmc7)
                  and row.get('monto') is not None and isinstance(row.get('fechaPago'), dict)
                  and row['fechaPago'].get('value')):
                eligible.append({'accion': 'DEPOSITAR', 'cheque': row})
        if eligible and verified_account:
            selected = eligible[0]
            cheque = selected['cheque']
            echeq = {}
            if cheque.get('chequeId'):
                echeq['idCheque'] = cheque['chequeId']
            if cheque.get('cmc7completo'):
                echeq['cmc7'] = cheque['cmc7completo']
            if selected['accion'] == 'DEPOSITAR':
                echeq['monto'] = str(cheque['monto'])
                echeq['fechaPago'] = cheque['fechaPago']['value'].split('T')[0]
            management_payload = {
                'idOrigen': str(uuid4()),
                'cbuCuenta': verified_account.get('CBU', verified_account.get('cbu')),
                'accion': selected['accion'],
                'echeqs': [echeq],
            }
            preview = local_request(
                api, recorder, 'POST', '/echeqs/gestion/previsualizar',
                body=management_payload,
                function=f'Previsualizar gestión eCheq {selected["accion"]} - NO ENVIADA',
            )
            results.append(('preview_gestion', preview.status_code if preview else None))
            details = {
                'accion': selected['accion'],
                'idCheque': cheque.get('chequeId'),
                'cmc7': cheque.get('cmc7completo'),
                'estado_banco': cheque.get('estado'),
                'monto_banco': cheque.get('monto'),
                'fechaPago_banco': cheque.get('fechaPago'),
                'payload_banco': (response_json(preview) or {}).get('payload_banco') if preview else None,
            }
            recorder.note(
                f'Gestión eCheq candidata {selected["accion"]}',
                'PENDIENTE_CONFIRMACION',
                'Cheque recibido devuelto por la API y estado compatible según la documentación. No se ejecutó '
                'la acción; requiere confirmación conjunta del cheque y payload.',
                identifiers={'idCheque': cheque.get('chequeId'), 'cmc7': cmc7},
            )
            notes.append({'gestion_candidata_no_ejecutada': details})
        else:
            received = lists.get('RECIBIDOS')
            if received is None or received.status_code != 200:
                reason = (
                    'La consulta de eCheqs RECIBIDOS no terminó con HTTP 200. No se infiere que no haya cheques. '
                    'Revisar el intercambio bancario, el scope y la conectividad.'
                )
                result = 'PRUEBA PENDIENTE'
            elif not eligible:
                reason = (
                    'La consulta RECIBIDOS no devolvió un cheque con estado EMITIDO-PENDIENTE o ACTIVO, '
                    'identificador real y datos requeridos por la acción. Solicitar al banco un eCheq de '
                    'homologación recibido por el adherente, su estado apto para ACEPTAR o DEPOSITAR y, '
                    'para depósito, monto y fechaPago informados por la API.'
                )
                result = 'PENDIENTE POR FALTA DE ECHEQS DE PRUEBA'
            else:
                reason = (
                    'Hay un cheque elegible, pero no se encontró una cuenta con CBU devuelta por la consulta '
                    'de cuentas para preparar la acción.'
                )
                result = 'PRUEBA PENDIENTE'
            recorder.note('Gestión de eCheqs', result, reason)
            notes.append(reason)

        write_preview(recorder, payloads, notes)

    evidence_dir = recorder.directory
    print(f'Evidencias: {evidence_dir.resolve() if evidence_dir else "(sin llamadas registradas)"}')
    print('Resultados HTTP locales:', json.dumps(results, ensure_ascii=False))
    print('PREVIAS_OPERACIONES:', json.dumps(payloads, ensure_ascii=False, indent=2))
    added = [(name, code) for name, code in results if name.startswith('alta:') and code == 200]
    if added:
        print(f'Altas condicionales confirmadas por ausencia bancaria: {len(added)}; verifique alta y reconsulta en evidencias.')
    else:
        print('No se hicieron altas condicionales en esta ejecución.')
    print('No se enviaron transferencias, emisiones de eCheq ni gestiones.')
    required_reads = [
        code for name, code in results
        if name == 'cuentas' or name.startswith('movimientos:') or name.startswith('lista_echeqs:')
    ]
    return 0 if required_reads and all(code == 200 for code in required_reads) else 1


if __name__ == '__main__':
    raise SystemExit(run())
